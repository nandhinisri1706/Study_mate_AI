from pptx import Presentation


def extract_text_from_ppt(file):

    presentation = Presentation(file)

    text = ""

    for slide in presentation.slides:

        for shape in slide.shapes:

            if hasattr(shape, "text"):
                text += shape.text + "\n"

    return text